"""Load the case-study workbook into plain Python structures.

The workbook stands in for three systems:
  * Fabric Catalog        -> ERP item master + inventory by dye lot
  * Pricing Rules         -> ERP price rules (parsed into PricingConfig)
  * Existing CRM Records  -> CRM accounts + contacts
  * Inbound Requests      -> website form submissions + inbound emails

Website requests are enriched from data/web_cart_submissions.json: on the real site a
sample is added to the cart from a product page (exact colorway SKU + Project Type), so
the website hands us structured lines, with the free text only as an order note.
"""
from __future__ import annotations

import json
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import openpyxl

DEFAULT_DATA = Path(__file__).resolve().parent.parent / "data" / "perennials_sample_data.xlsx"
CARTS_FILE = "web_cart_submissions.json"  # looked up next to the workbook


@dataclass
class PricingConfig:
    trade_discount: float = 0.15
    non_trade_discount: float = 0.0
    volume_tiers: tuple = ((100, 0.10), (50, 0.05))  # (min yards, extra discount), highest first
    cut_fee: float = 25.0
    cut_fee_under_yards: float = 5.0
    max_samples: int = 10
    quote_valid_days: int = 30
    hold_days: int = 14


def _pct(s) -> float:
    return float(re.sub(r"[^\d.]", "", str(s)) or 0) / 100


def _num(s) -> float:
    return float(re.sub(r"[^\d.]", "", str(s)) or 0)


def _rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() for h in rows[0]]
    return [dict(zip(header, r)) for r in rows[1:] if any(v is not None for v in r)]


def load_workbook(path: str | Path = DEFAULT_DATA) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)

    lots = []
    for r in _rows(wb["Fabric Catalog"]):
        lots.append({
            "sku": r["SKU"], "fabric": r["Fabric"], "color": r["Color"], "status": r["Status"],
            "list_price": float(r["List Price ($/yd)"]), "lot": r["Dye Lot"],
            "on_hand": float(r["Yards On Hand"]), "reserved": float(r["Yards Reserved"]),
        })

    # Pricing rules: parse the sheet so a merchandiser can change numbers in Excel.
    cfg = PricingConfig()
    tiers = []
    for r in _rows(wb["Pricing Rules"]):
        rule, val = str(r["Rule"]).lower(), r["Value"]
        if rule.startswith("trade discount"):
            cfg.trade_discount = _pct(val)
        elif rule.startswith("non-trade"):
            cfg.non_trade_discount = _pct(val)
        elif rule.startswith("volume tier"):
            m = re.search(r"(\d+)", rule)
            tiers.append((int(m.group(1)), _pct(val)))
        elif rule.startswith("cut fee"):
            cfg.cut_fee = _num(val)
            m = re.search(r"under (\d+)", str(r.get("Notes") or ""))
            if m:
                cfg.cut_fee_under_yards = float(m.group(1))
        elif rule.startswith("samples"):
            m = re.search(r"max (\d+)", str(r.get("Notes") or ""))
            if m:
                cfg.max_samples = int(m.group(1))
        elif rule.startswith("quote validity"):
            cfg.quote_valid_days = int(_num(val))
        elif rule.startswith("reservation hold"):
            cfg.hold_days = int(_num(val))
    if tiers:
        cfg.volume_tiers = tuple(sorted(tiers, reverse=True))

    accounts, contacts = [], []
    for r in _rows(wb["Existing CRM Records"]):
        if r["Record Type"] == "Account":
            accounts.append({"id": r["ID"], "name": r["Name"], "type": r["Type"], "status": r["Status"]})
        elif r["Record Type"] == "Contact":
            contacts.append({"id": r["ID"], "account_id": r["Account ID"], "name": r["Name"],
                             "email": (r["Email"] or "").lower(), "alt_emails": []})

    requests = []
    for r in _rows(wb["Inbound Requests"]):
        requests.append({
            "id": r["ID"], "type": r["Type"], "channel": r["Channel"],
            "received": datetime.strptime(str(r["Received (UTC)"]), "%Y-%m-%d %H:%M"),
            "name": r["Name"], "company": r["Company"], "email": (r["Email"] or "").lower(),
            "message": r["Message"] or "",
        })
    requests.sort(key=lambda x: x["received"])

    carts_path = Path(path).with_name(CARTS_FILE)
    carts = json.loads(carts_path.read_text()) if carts_path.exists() else {}
    for r in requests:
        if r["id"] in carts:
            r["cart"] = carts[r["id"]]

    return {"lots": lots, "pricing": cfg, "accounts": accounts, "contacts": contacts, "requests": requests}
